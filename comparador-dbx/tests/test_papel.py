"""L5 · Papel de trabajo (docs/10, docs/13): workpaper, annex y styles. Sin modelo.

Las marcas de cada artículo vienen de un GUION (`datos_papel.entrada`) hasta que exista L4.
"""
import ast
import dataclasses
import re
import socket
import subprocess
import sys
import zipfile
from collections import Counter
from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.config import settings
from backend.output import annex, styles, workpaper
from backend.output.workpaper import generar_papel
from datos_papel import entrada, par, seccion, veredicto

COLUMNAS_PAPEL = ["Nombre de la normativa", "Sección / Artículo", "Manual interno", "Verificación",
                  "Comentario", "Referencias / Evidencias"]
COLUMNAS_ANEXO = ["ID par", "Artículo", "Sección manual", "Origen", "Similitud v1", "Similitud v2",
                  "Similitud máxima", "Rango v1", "Rango v2", "Naturaleza", "Cobertura", "Confianza",
                  "Elegido para el papel", "Comentario del juez", "Banderas"]
H = 12            # fila del encabezado de la matriz (10 etiquetas + 1 en blanco + 1)
HOSTILES = ["=1+1", "+cmd", "-2+3", "@SUM(A1)", "\t=1", "\r=1", " =1+1", '=HYPERLINK("http://x.test","y")']


def guardar(ent, tmp_path):
    ruta = tmp_path / "papel.xlsx"
    conteos = generar_papel(ent, ruta)
    return conteos, load_workbook(ruta), ruta


@pytest.fixture
def libro(tmp_path):
    return guardar(entrada(), tmp_path)


def fila_matriz(ws, n):
    return [c.value for c in ws[n]][:6]


# ---------- Hoja 1: encabezado, leyenda, matriz ----------

def test_dos_hojas_y_sin_tercera(libro):
    assert libro[1].sheetnames == ["Papel de trabajo", "Anexo técnico"]


def test_encabezado_etiquetas_y_valores(libro):
    ws = libro[1]["Papel de trabajo"]
    assert [ws.cell(row=i, column=1).value for i in range(1, 11)] == workpaper.ETIQUETAS
    assert ws["B1"].value in (None, "")                       # revisión en blanco
    assert ws["B4"].value in (None, "")                       # corte en blanco
    assert ws["B2"].value == "Evaluación de cumplimiento normativo — Manual Interno.pdf"
    assert ws["B3"].value == workpaper.OBJETIVO
    assert ws["B8"].value in (None, "") and ws["B9"].value in (None, "")
    assert "Norma Uno.pdf" in ws["B5"].value and "versus" in ws["B5"].value and "Manual Interno.pdf" in ws["B5"].value
    assert ws["B10"].value == date.today().strftime("%d/%m/%Y")


def test_fecha_de_ejecucion_dada(tmp_path):
    ws = guardar(entrada(fecha_ejecucion=date(2026, 10, 7)), tmp_path)[1]["Papel de trabajo"]
    assert ws["B10"].value == "07/10/2026"


def test_leyenda_fija_en_orden(libro):
    leyenda = libro[1]["Papel de trabajo"]["B6"].value
    assert leyenda == ("A = Sí cumple  ·  L = Parcialmente  ·  R = No cumple  ·  X = No aplica  ·  "
                       "P = Información obtenida de la normativa")


@pytest.mark.parametrize("conclusion,editada,esperado", [
    ("Texto.", False, workpaper.PREFIJO_BORRADOR + "\nTexto."),
    ("Texto.", True, "Texto."),
    (workpaper.PREFIJO_BORRADOR + "\nYa lleva prefijo", False, workpaper.PREFIJO_BORRADOR + "\nYa lleva prefijo"),
    ("", False, None),
])
def test_conclusion_con_prefijo_de_borrador(tmp_path, conclusion, editada, esperado):
    ws = guardar(entrada(conclusion=conclusion, conclusion_editada=editada), tmp_path)[1]["Papel de trabajo"]
    assert (ws["B7"].value or None) == esperado


def test_columnas_de_la_matriz_en_orden_exacto(libro):
    ws = libro[1]["Papel de trabajo"]
    assert [ws.cell(row=H, column=c).value for c in range(1, 7)] == COLUMNAS_PAPEL
    assert ws.cell(row=H, column=7).value is None
    assert [ws.column_dimensions[l].width for l in "ABCDEF"] == [18, 45, 45, 10, 40, 28]


def test_orden_de_filas_es_el_del_documento_normativo(libro):
    ws = libro[1]["Papel de trabajo"]
    marcas = [ws.cell(row=r, column=4).value for r in range(H + 1, H + 7)]
    assert marcas == ["A", "L", "R", "X", "P", "A"]            # Art.1..5 de N1 y luego Art.9 de N2
    assert ws.cell(row=H + 7, column=4).value is None


def test_paneles_congelados_bajo_el_encabezado(libro):
    assert libro[1]["Papel de trabajo"].freeze_panes == f"A{H + 1}"
    assert libro[1]["Anexo técnico"].freeze_panes == "A2"


def test_encabezado_con_colores_de_settings(libro):
    c = libro[1]["Papel de trabajo"].cell(row=H, column=1)
    assert c.fill.start_color.rgb.endswith(settings.COLOR_PRIMARIO)
    assert c.font.color.rgb.endswith(settings.COLOR_FONDO) and c.font.bold


def test_colores_siguen_a_settings(monkeypatch, tmp_path):
    monkeypatch.setattr(styles, "settings", dataclasses.replace(settings, COLOR_PRIMARIO="112233", COLOR_FONDO="FEDCBA"))
    wb = guardar(entrada(), tmp_path)[1]
    for hoja, fila in (("Papel de trabajo", H), ("Anexo técnico", 1)):
        c = wb[hoja].cell(row=fila, column=1)
        assert c.fill.start_color.rgb.endswith("112233") and c.font.color.rgb.endswith("FEDCBA")


def test_color_suave_por_marca_y_centrada(libro):
    ws = libro[1]["Papel de trabajo"]
    for r, marca in zip(range(H + 1, H + 7), "ALRXPA"):
        c = ws.cell(row=r, column=4)
        assert c.value == marca and c.fill.start_color.rgb.endswith(styles.COLOR_MARCA[marca])
        assert c.alignment.horizontal == "center"


def test_celdas_con_ajuste_y_alineacion_superior(libro):
    c = libro[1]["Papel de trabajo"].cell(row=H + 1, column=2)
    assert c.alignment.wrap_text and c.alignment.vertical == "top"


def test_col1_combinada_verticalmente_por_normativa(libro):
    ws = libro[1]["Papel de trabajo"]
    combinadas = sorted(str(m) for m in ws.merged_cells.ranges if m.min_col == 1 and m.min_row > H)
    assert combinadas == [f"A{H + 1}:A{H + 5}"]               # N1 (5 artículos); N2 tiene uno: sin combinar
    assert ws.cell(row=H + 1, column=1).value == "Norma Uno.pdf\nLibro I › Título II"
    assert ws.cell(row=H + 6, column=1).value == "Norma Dos.pdf\nLibro IX"


def test_col2_ruta_con_separador_y_texto_literal(tmp_path):
    ent = entrada()
    ent.filas[3].articulo.texto_literal = "Literal del art. 1\nsegunda línea"
    ws = guardar(ent, tmp_path)[1]["Papel de trabajo"]
    assert ws.cell(row=H + 1, column=2).value == "Capítulo III › Art. 1\nLiteral del art. 1\nsegunda línea"
    assert ws.cell(row=H + 2, column=2).value.startswith("Capítulo III › Art. 2 Reporte\n")


def test_col3_vacia_en_R_X_P_y_llena_en_A_L(libro):
    ws = libro[1]["Papel de trabajo"]
    col3 = {ws.cell(row=r, column=4).value + str(r): ws.cell(row=r, column=3).value for r in range(H + 1, H + 7)}
    assert col3[f"A{H + 1}"] == "Capítulo 3 › 3.1 Conciliación\ncita del manual"
    assert col3[f"L{H + 2}"] == "Capítulo 3 › 3.2\ncita del manual"
    for r in (H + 3, H + 4, H + 5):
        assert ws.cell(row=r, column=3).value in (None, "")
    assert ws.cell(row=H + 6, column=3).value.startswith("Capítulo 4 › 4.1")


def test_comentario_y_elementos_faltantes(libro):
    ws = libro[1]["Papel de trabajo"]
    assert ws.cell(row=H + 3, column=5).value == "Brecha.\nElementos faltantes:\n- plazo\n- responsable"
    assert ws.cell(row=H + 1, column=5).value == "Cumple a1."


def test_referencias_archivo_pagina_similitud_y_avisos(libro):
    ws = libro[1]["Papel de trabajo"]
    a = ws.cell(row=H + 1, column=6).value
    assert a == "Norma: Norma Uno.pdf, p. 1\nManual: Manual Interno.pdf, sección 3.1, p. 7\nSimilitud: 0.840"
    l = ws.cell(row=H + 2, column=6).value                    # cita del manual no verificada + sección incierta
    assert "pp. " not in l.split("\n")[0] and "Manual: Manual Interno.pdf, sección 3.2, pp. 8-9" in l
    assert "Aviso: la cita no se verificó contra el texto fuente" in l
    assert "Aviso: el límite del artículo es incierto" in l
    r = ws.cell(row=H + 3, column=6).value                    # R: sin manual ni similitud, pero con aviso de cita
    assert "Manual:" not in r and "Similitud" not in r and "Aviso: la cita no se verificó" in r
    assert ws.cell(row=H + 5, column=6).value == "Norma: Norma Uno.pdf, pp. 5-6"


# ---------- Hoja 2: columnas, 'Elegido', banderas, vía 2 ----------

def filas_anexo(wb):
    ws = wb["Anexo técnico"]
    return [[c.value for c in ws[r]] for r in range(2, 13)]


def test_columnas_del_anexo_en_orden_exacto(libro):
    ws = libro[1]["Anexo técnico"]
    assert [ws.cell(row=1, column=c).value for c in range(1, 16)] == COLUMNAS_ANEXO
    assert ws.cell(row=1, column=16).value is None
    assert annex.COLUMNAS == COLUMNAS_ANEXO


def test_anexo_registra_todos_los_pares(libro):
    assert [f[0] for f in filas_anexo(libro[1])] == [f"P{i}" for i in range(1, 12)]


def test_elegido_para_el_papel_si_o_no_con_motivo(libro):
    esperado = {"P1": "Sí", "P2": "No: cobertura menor", "P3": "No: confianza menor",
                "P4": "No: sin veredicto: requiere revisión", "P5": "Sí", "P6": "No: confianza menor",
                "P7": "No: cobertura menor", "P8": "No: cobertura nula: ningún par cubre la obligación",
                "P9": "No: no aplica a la entidad: sin respaldo", "P10": "No: artículo informativo: sin respaldo",
                "P11": "Sí"}
    assert {f[0]: f[12] for f in filas_anexo(libro[1])} == esperado


def test_motivos_de_empate_y_sin_respaldo():
    a = par("Pa", "x", "y", veredicto())
    b = par("Pb", "x", "z", veredicto())
    assert annex._motivo(b, a, "A") == "empate: se eligió otro par"
    assert annex._motivo(b, None, "A") == "no se eligió respaldo"


def test_celdas_de_articulo_y_seccion_y_origen(libro):
    f = {x[0]: x for x in filas_anexo(libro[1])}
    assert f["P1"][1] == "N1-a1 | Art. 1 | Norma Uno.pdf" and f["P1"][2] == "M1-s1 | 3.1 | Manual Interno.pdf"
    assert (f["P1"][3], f["P2"][3], f["P3"][3]) == ("ambas", "v1", "v2")
    assert f["P1"][12] == "Sí" and f["P1"][13] == "Comentario del juez."


def test_similitudes_rangos_y_vacios(libro):
    f = {x[0]: x for x in filas_anexo(libro[1])}
    assert f["P1"][4:9] == [0.81, 0.84, 0.84, 1, 2]
    assert f["P3"][4:9] == [None, 0.6, 0.6, None, 5]
    assert f["P4"][4:9] == [None] * 5 and f["P4"][9:12] == [None, None, None]   # sin veredicto ni scores
    assert f["P1"][9:12] == ["obligacion", "total", 0.9]


def test_banderas(libro):
    f = {x[0]: x[14] for x in filas_anexo(libro[1])}
    assert f["P1"] in (None, "") and f["P3"] in (None, "")
    assert f["P5"] == "cita_no_verificada, seccionado_incierto"
    assert f["P6"] == "requiere_revision, seccionado_incierto"
    assert f["P8"] == "cita_no_verificada"


def test_bloque_via2_al_final_de_la_hoja_2(libro):
    ws = libro[1]["Anexo técnico"]
    r = 11 + 4                                                # 11 pares + 4
    assert ws.cell(row=r, column=1).value == annex.TITULO_VIA2
    assert f"A{r}:C{r}" in [str(m) for m in ws.merged_cells.ranges]
    assert [ws.cell(row=r + 1, column=c).value for c in range(1, 7)] == annex.COLUMNAS_VIA2
    assert [ws.cell(row=r + 2, column=c).value for c in range(1, 7)] == \
        ["M1-c1", "9.1", "Manual Interno.pdf", "Capítulo 9", "15", "Control huérfano 1"]
    assert ws.cell(row=r + 3, column=5).value == "16-17"
    assert ws.max_row == r + 3 and len(libro[1].sheetnames) == 2


def test_via2_vacia_dice_ninguno(tmp_path):
    ws = guardar(entrada(controles_sin_base=[]), tmp_path)[1]["Anexo técnico"]
    assert ws.cell(row=15, column=1).value == annex.TITULO_VIA2 and ws.cell(row=17, column=1).value == "(ninguno)"


def test_pares_de_articulos_desconocidos_no_rompen(tmp_path):
    ent = entrada()
    ent.pares.append(par("PX", "N9-zz", "M1-s1", veredicto()))
    assert [f[0] for f in filas_anexo(guardar(ent, tmp_path)[1])][-1] == "P11"


# ---------- Conteos Hoja 1 = derivados de Hoja 2 (docs/09 §5) ----------

def marca_desde_hoja2(filas):
    """docs/09 §5 aplicado SOLO a lo que dice la Hoja 2 de un artículo (naturaleza y cobertura)."""
    validos = [f for f in filas if f[9]]
    cuenta = Counter(f[9] for f in validos).most_common()
    nat = "obligacion" if len(cuenta) > 1 and cuenta[0][1] == cuenta[1][1] else cuenta[0][0]
    if nat == "informativo":
        return "P"
    if nat == "no_aplica_entidad":
        return "X"
    coberturas = {f[10] for f in validos}
    return "A" if "total" in coberturas else "L" if "parcial" in coberturas else "R"


def test_conteos_hoja1_igual_a_derivados_de_hoja2(libro):
    conteos, wb, _ = libro
    por_articulo = {}
    for f in filas_anexo(wb):
        por_articulo.setdefault(f[1].split(" | ")[0], []).append(f)
    derivados = Counter(marca_desde_hoja2(v) for v in por_articulo.values())
    ws = wb["Papel de trabajo"]
    hoja1 = Counter(ws.cell(row=r, column=4).value for r in range(H + 1, H + 7))
    assert conteos == {m: derivados.get(m, 0) for m in "ALRXP"} == {m: hoja1.get(m, 0) for m in "ALRXP"}
    assert conteos == {"A": 2, "L": 1, "R": 1, "X": 1, "P": 1}
    # Solo A y L tienen un par "Sí" en la Hoja 2
    assert sum(1 for f in filas_anexo(wb) if f[12] == "Sí") == conteos["A"] + conteos["L"]


# ---------- styles: neutralización, truncado, NaN/inf ----------

@pytest.mark.parametrize("texto", HOSTILES)
def test_escribir_neutraliza_formulas(tmp_path, texto):
    from openpyxl import Workbook
    ws = Workbook().active
    c = styles.escribir(ws, 1, 1, texto)
    assert c.data_type == "s" and c.value == texto             # texto intacto, nunca fórmula
    assert c.quotePrefix == texto.startswith(styles.PREFIJOS_FORMULA)


@pytest.mark.parametrize("texto", HOSTILES)
def test_papel_con_texto_hostil_no_tiene_formulas(tmp_path, texto):
    """El texto hostil entra por todas las puertas de texto libre: comentario, juez, vía 2, conclusión."""
    ent = entrada(conclusion=texto, conclusion_editada=True)
    ent.filas[3].comentario = texto
    ent.pares[0].veredicto.comentario = texto
    ent.controles_sin_base[0].texto_literal = texto
    ent.filas[3].articulo.texto_literal = texto
    _, wb, ruta = guardar(ent, tmp_path)
    for ws in wb:
        for fila in ws.iter_rows():
            for c in fila:
                assert c.data_type != "f", (ws.title, c.coordinate, c.value)
    with zipfile.ZipFile(ruta) as z:
        for nombre in z.namelist():
            if nombre.startswith("xl/worksheets/"):
                assert not re.search(r"<f[ >/]", z.read(nombre).decode("utf-8")), nombre
    # openpyxl no conserva los espacios/tabuladores iniciales al releer: se compara sin ellos
    assert wb["Papel de trabajo"]["B7"].value.strip() == texto.strip()
    assert wb["Papel de trabajo"].cell(row=H + 1, column=5).value.strip() == texto.strip()


def test_truncado_a_32767_con_aviso_de_pagina():
    largo = "x" * 40000
    t = styles.texto_seguro(largo, pagina=7)
    assert len(t) == styles.MAX_CELDA == 32767 and t.endswith("[texto truncado, ver página 7]")
    assert styles.texto_seguro(largo).endswith("[texto truncado]")
    assert len(styles.texto_seguro("x" * 32767, 1)) == 32767 and "truncado" not in styles.texto_seguro("x" * 32767, 1)
    assert len(styles.texto_seguro("x" * 32768, 1)) == 32767


def test_truncado_en_el_papel_ya_guardado(tmp_path):
    ent = entrada()
    ent.filas[3].articulo.texto_literal = "y" * 50000
    ws = guardar(ent, tmp_path)[1]["Papel de trabajo"]
    c = ws.cell(row=H + 1, column=2).value
    assert len(c) <= 32767 and c.endswith("[texto truncado, ver página 1]")


def test_caracteres_ilegales_se_quitan():
    assert styles.texto_seguro("a\x00b\x07c") == "abc"


@pytest.mark.parametrize("valor", [float("nan"), float("inf"), float("-inf")])
def test_numeros_no_finitos_quedan_vacios(valor):
    from openpyxl import Workbook
    c = styles.escribir(Workbook().active, 1, 1, valor)
    assert c.value is None


def test_numeros_y_nulos_normales():
    from openpyxl import Workbook
    ws = Workbook().active
    assert styles.escribir(ws, 1, 1, 0.5).value == 0.5 and styles.escribir(ws, 1, 2, 3).value == 3
    assert styles.escribir(ws, 1, 3, None).value in (None, "")


def test_nan_e_inf_en_el_anexo_quedan_vacios(tmp_path):
    ent = entrada()
    ent.pares[1].score_v1 = float("nan")        # P2 (solo v1): sin ningún score válido
    ent.pares[1].score_v2 = float("inf")
    ws = guardar(ent, tmp_path)[1]["Anexo técnico"]
    assert [ws.cell(row=3, column=c).value for c in (5, 6)] == [None, None]


def test_similitud_maxima_ignora_scores_no_finitos_qa02(tmp_path):
    ent = entrada()
    ent.pares[0].score_v1 = float("nan")        # P1 tiene v1 = nan y v2 = 0.84: el máximo debe ser 0.84
    libro = guardar(ent, tmp_path)[1]
    assert libro["Anexo técnico"].cell(row=2, column=7).value == 0.84
    assert "Similitud: 0.840" in libro["Papel de trabajo"].cell(row=H + 1, column=6).value


def test_alto_de_fila_tiene_tope():
    assert styles.alto_fila([("a\n" * 1000, 10)]) == styles.MAX_ALTO
    assert styles.alto_fila([("corto", 40)]) == 15


# ---------- RNF-08: cero llamadas al modelo ----------

def test_modulos_de_papel_no_importan_backend_llm():
    for modulo in (workpaper, annex, styles):
        arbol = ast.parse(Path(modulo.__file__).read_text(encoding="utf-8"))
        importados = [n.module if isinstance(n, ast.ImportFrom) else a.name
                      for n in ast.walk(arbol) if isinstance(n, (ast.Import, ast.ImportFrom))
                      for a in (n.names if isinstance(n, ast.Import) else [n])]
        assert not [i for i in importados if i and i.startswith("backend.llm")], modulo.__name__


def test_importar_el_papel_no_carga_backend_llm():
    codigo = "import sys, backend.output.workpaper; print([m for m in sys.modules if m.startswith('backend.llm')])"
    raiz = Path(__file__).resolve().parent.parent
    salida = subprocess.run([sys.executable, "-c", codigo], cwd=raiz, capture_output=True, text=True, check=True)
    assert salida.stdout.strip() == "[]"


def test_generar_papel_no_abre_conexiones(monkeypatch, tmp_path):
    def prohibido(*a, **k):
        raise AssertionError("generar_papel intentó abrir una conexión")
    monkeypatch.setattr(socket.socket, "connect", prohibido)
    monkeypatch.setattr(socket, "create_connection", prohibido)
    monkeypatch.setattr(socket, "getaddrinfo", prohibido)
    guardar(entrada(), tmp_path)


def test_generar_papel_acepta_archivo_en_memoria():
    buf = BytesIO()
    assert generar_papel(entrada(), buf) == {"A": 2, "L": 1, "R": 1, "X": 1, "P": 1}
    assert load_workbook(BytesIO(buf.getvalue())).sheetnames == ["Papel de trabajo", "Anexo técnico"]


def test_papel_sin_filas_no_falla(tmp_path):
    conteos, wb, _ = guardar(entrada(filas=[], pares=[], controles_sin_base=[]), tmp_path)
    assert conteos == dict.fromkeys("ALRXP", 0) and wb["Papel de trabajo"].max_row == H
