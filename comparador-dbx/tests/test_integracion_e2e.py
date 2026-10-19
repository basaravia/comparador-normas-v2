"""Extremo a extremo con los modelos REALES del stage (LLM y embeddings) y FAISS real, sobre el ejemplo precargado. Marca `integracion`: gasta tokens."""
import time

import pytest
from openpyxl import load_workbook

from backend.core import sesiones
from backend.engine import servicio

pytestmark = pytest.mark.integracion


def _hojas(nodos):
    for n in nodos:
        if n.hoja:
            yield n
        yield from _hojas(n.hijos)


def test_ejemplo_a_excel_con_modelos_reales(cliente):
    s = sesiones.crear()
    try:
        servicio.cargar_ejemplo(s)
        arts = [h.id for h in _hojas(servicio.arbol(s, "normativa")[0].nodos)][:1]
        secs = [h.id for h in _hojas(servicio.arbol(s, "manual_control")[0].nodos) if h.seleccionada][:3]
        job = servicio.comparar(s, arts, secs, cliente)
        t0 = time.time()
        while job.estado in ("en_cola", "en_curso") and time.time() - t0 < 900:
            time.sleep(2)
        assert job.estado == "listo", job.errores
        comp = s.comparaciones[job.comparacion_id]
        assert comp.out.pares_juzgados >= 1 and len(comp.out.via1) == 1 and comp.out.via1[0].marca in "ALRXP"
        for f in comp.out.via1:                                   # toda cita que llega al resultado es subcadena del texto fuente
            norm = lambda t: " ".join(t.split())
            if f.cita_manual:
                fuente = norm(servicio.seccion(s, f.seccion_respaldo_id).texto)
                assert norm(f.cita_manual) in fuente or f.requiere_revision
        ruta = servicio.excel(s, comp)
        wb = load_workbook(ruta)
        assert wb.sheetnames == ["Papel de trabajo", "Anexo técnico"]
    finally:
        sesiones.reiniciar(s.id)
