"""Estilos del papel de trabajo y la única puerta de entrada de texto a las celdas.

Todo texto que llega al Excel pasa por `escribir`: limpia caracteres ilegales, trunca a
32.767 caracteres y neutraliza el prefijo de fórmula (CWE-1236, OWASP LLM05).
"""
import math
import re

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from backend.config import settings

MAX_CELDA = 32767                                    # límite de Excel por celda
PREFIJOS_FORMULA = ("=", "+", "-", "@", "\t", "\r")   # Excel los interpreta como fórmula (OWASP)
ANCHOS_MATRIZ = [18, 45, 45, 10, 40, 28]              # docs/10: anchos orientativos
MAX_ALTO = 409                                        # alto máximo de fila en Excel (puntos)

# Colores suaves por marca (docs/10); el verde institucional viene de la configuración.
COLOR_MARCA = {"A": "C6EFCE", "L": "FFEB9C", "R": "F8CBAD", "X": "E7E6E6", "P": "DDEBF7"}
SIGNIFICADO_MARCA = {"A": "Sí cumple", "L": "Parcialmente", "R": "No cumple",
                     "X": "No aplica", "P": "Información obtenida de la normativa"}

_borde = Side(style="thin", color="BFBFBF")
BORDE = Border(left=_borde, right=_borde, top=_borde, bottom=_borde)
ARRIBA = Alignment(wrap_text=True, vertical="top")
CENTRADO = Alignment(wrap_text=True, vertical="top", horizontal="center")


def relleno(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


def fuente_encabezado() -> Font:
    return Font(bold=True, size=9, color=settings.COLOR_FONDO)    # blanco sobre el verde


def relleno_encabezado() -> PatternFill:
    return relleno(settings.COLOR_PRIMARIO)


def texto_seguro(valor, pagina: int | None = None) -> str:
    """Texto listo para una celda: sin caracteres ilegales y con tope de 32.767 caracteres."""
    texto = ILLEGAL_CHARACTERS_RE.sub("", str(valor))
    if len(texto) <= MAX_CELDA:
        return texto
    aviso = f"[texto truncado, ver página {pagina}]" if pagina else "[texto truncado]"
    return texto[: MAX_CELDA - len(aviso) - 1] + "\n" + aviso


def escribir(ws, fila: int, col: int, valor, pagina: int | None = None, estilo=None):
    """Escribe una celda. Los números van como número (los no finitos, como celda vacía); todo lo demás, como TEXTO.

    Si el texto empieza por `=`, `+`, `-`, `@`, tabulador o retorno de carro, se fuerza el tipo
    cadena y `quotePrefix` (Excel no lo evalúa ni al editar la celda). El texto no se altera.
    """
    celda = ws.cell(row=fila, column=col)
    if isinstance(valor, float) and not math.isfinite(valor):
        celda.value = None                         # NaN e inf: celda vacía (Excel no los admite)
    elif isinstance(valor, (int, float)) and not isinstance(valor, bool):
        celda.value = valor
    else:
        celda.value = texto_seguro(valor if valor is not None else "", pagina)
        celda.data_type = "s"                      # nunca "f", aunque empiece por "="
        if celda.value.startswith(PREFIJOS_FORMULA):
            celda.quotePrefix = True
    celda.alignment = ARRIBA
    celda.border = BORDE
    if estilo:
        estilo(celda)
    return celda


def alto_fila(textos_y_anchos: list[tuple[str, int]]) -> float:
    """Alto aproximado (puntos) para texto con ajuste: Excel no autoajusta celdas combinadas."""
    lineas = 1
    for texto, ancho in textos_y_anchos:
        n = sum(max(1, math.ceil(len(t) / max(ancho - 2, 1))) for t in str(texto).split("\n"))
        lineas = max(lineas, n)
    return min(MAX_ALTO, 15 * lineas)


def limpiar_nombre(texto: str) -> str:
    return re.sub(r"\s+", " ", texto).strip()


# Estilos reutilizables (reciben la celda)
def estilo_etiqueta(celda) -> None:
    celda.font = Font(bold=True)


def estilo_encabezado(celda) -> None:
    celda.font = fuente_encabezado()
    celda.fill = relleno_encabezado()
    celda.alignment = CENTRADO


def estilo_marca(marca: str):
    def aplicar(celda) -> None:
        celda.fill = relleno(COLOR_MARCA[marca])
        celda.alignment = CENTRADO
        celda.font = Font(bold=True)
    return aplicar


def ajustar_impresion(ws, fila_titulos: int) -> None:
    """Horizontal, ancho de una página y encabezado de la tabla repetido en cada página."""
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = f"{fila_titulos}:{fila_titulos}"
