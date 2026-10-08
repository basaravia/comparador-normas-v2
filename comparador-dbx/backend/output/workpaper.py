"""Papel de trabajo en Excel (docs/10): Hoja 1 'Papel de trabajo' y Hoja 2 'Anexo técnico'.

100 % Python con openpyxl (el anexo se arma con pandas). Cero llamadas al modelo: este módulo
no importa nada de `backend.llm`. Origen: módulo nuevo según docs/10 (la v6 no genera Excel).

La marca de cada artículo (A/L/R/X/P) llega YA calculada en `FilaPapel`: la agregación es de L4
(`backend/engine/aggregate.py`, docs/09 §5) y no se repite aquí.
"""
import math
from collections import Counter
from datetime import date
from typing import Literal, Optional

from openpyxl import Workbook
from pydantic import BaseModel

from backend.models.schemas import Documento, Par, Seccion
from backend.output import annex
from backend.output.styles import (ANCHOS_MATRIZ, BORDE, SIGNIFICADO_MARCA, ajustar_impresion, alto_fila,
                                   escribir, estilo_encabezado, estilo_etiqueta, estilo_marca)

PREFIJO_BORRADOR = "[Borrador generado automáticamente — validar]"
OBJETIVO = ("Evaluar el cumplimiento de la normativa regulatoria aplicable por parte del manual "
            "de control interno de la entidad, identificando brechas y evidencia de respaldo en "
            "ambas vías (norma → manual y manual → norma).")
SEP = " › "
COLUMNAS = ["Nombre de la normativa", "Sección / Artículo", "Manual interno", "Verificación",
            "Comentario", "Referencias / Evidencias"]
ETIQUETAS = ["Nombre de la revisión", "Nombre del papel de trabajo", "Objetivo", "Corte o periodo",
             "Fuente", "Marcas de verificación", "Conclusiones", "Elaborado por", "Revisado por",
             "Fecha de ejecución"]
FILA_ENCABEZADO_MATRIZ = len(ETIQUETAS) + 2     # una fila en blanco entre encabezado y matriz


class FilaPapel(BaseModel):
    """Una fila del papel = un artículo evaluado, con su marca ya calculada (entrada de L4)."""
    articulo: Seccion
    marca: Literal["A", "L", "R", "X", "P"]
    par: Optional[Par] = None            # par de respaldo (1); None para X y P
    respaldo: Optional[Seccion] = None   # sección del manual de ese par (1)
    comentario: str = ""                 # razonamiento del juez (viene del veredicto)
    elementos_faltantes: list[str] = []


class SeccionExcluida(BaseModel):
    """Sección que no se comparó por no ser cuerpo del documento (carátula, índice, anexo…) y su motivo (RNF-06, docs/10)."""
    documento: str
    identificador: str
    titulo: str = ""
    ruta: str = ""
    paginas: str = ""
    motivo: str


def secciones_excluidas(tabla) -> list[SeccionExcluida]:
    """Convierte `seleccion.excluidas(tabla)` (DataFrame) en la lista que lleva `EntradaPapel`."""
    return [SeccionExcluida(documento=r["Documento"], identificador=r["Identificador"], titulo=r["Título"], ruta=r["Ruta"],
                            paginas=str(r["Págs"]), motivo=r["Motivo"]) for _, r in tabla.iterrows()]


class EntradaPapel(BaseModel):
    """Todo lo que necesita `generar_papel`. No incluye ningún objeto del modelo ni del índice."""
    manual: Documento
    normativas: list[Documento]
    filas: list[FilaPapel]                       # una por artículo (cualquier orden)
    pares: list[Par]                             # TODOS los pares evaluados (anexo)
    secciones: dict[str, Seccion]                # id → Seccion: artículos y secciones del manual
    controles_sin_base: list[Seccion] = []       # vía 2 (L4: controles_sin_base)
    excluidas: list[SeccionExcluida] = []        # no comparables, con su motivo (decisión del usuario, docs/14 #24)
    conclusion: str = ""                         # borrador (o texto editado) sin prefijo
    conclusion_editada: bool = False
    fecha_ejecucion: Optional[date] = None       # por defecto, hoy


def conteos_papel(entrada: EntradaPapel) -> dict[str, int]:
    c = Counter(f.marca for f in entrada.filas)
    return {m: c.get(m, 0) for m in "ALRXP"}


def _ordenar(entrada: EntradaPapel) -> list[FilaPapel]:
    """Orden del documento normativo: por documento (orden de `normativas`) y por posición."""
    orden_doc = {d.id: i for i, d in enumerate(entrada.normativas)}
    return sorted(entrada.filas, key=lambda f: (orden_doc.get(f.articulo.doc_id, 999),
                                               f.articulo.pagina_inicio, f.articulo.id))


def _fuente(entrada: EntradaPapel) -> str:
    """Libro, título y capítulo de cada normativa versus el manual."""
    lineas = []
    for d in entrada.normativas:
        rutas = []
        for f in entrada.filas:
            r = SEP.join(f.articulo.ruta[:3])
            if f.articulo.doc_id == d.id and r and r not in rutas:
                rutas.append(r)
        lineas.append(f"{d.nombre}: {'; '.join(rutas) or '—'}")
    return "\n".join(lineas) + f"\nversus\n{entrada.manual.nombre}"


def _nombre(s: Seccion) -> str:
    return s.identificador + (f" {s.titulo}" if s.titulo else "")


def _col1(f: FilaPapel, docs: dict[str, Documento]) -> str:
    nombre = docs[f.articulo.doc_id].nombre if f.articulo.doc_id in docs else f.articulo.doc_id
    return nombre + "\n" + SEP.join(f.articulo.ruta[:2])


def _col2(f: FilaPapel) -> str:
    a = f.articulo
    return SEP.join(a.ruta[2:] + [_nombre(a)]) + "\n" + a.texto_literal


def _con_respaldo(f: FilaPapel) -> bool:
    """Solo A y L muestran sección del manual; R, X y P la dejan vacía (docs/10)."""
    return f.marca in ("A", "L") and f.respaldo is not None


def _col3(f: FilaPapel) -> str:
    if not _con_respaldo(f):
        return ""
    cita = f.par.veredicto.cita_manual if f.par and f.par.veredicto and f.par.veredicto.cita_manual else ""
    return SEP.join(f.respaldo.ruta + [_nombre(f.respaldo)]) + ("\n" + cita if cita else "")


def _col5(f: FilaPapel) -> str:
    texto = f.comentario
    if f.elementos_faltantes:
        texto += "\nElementos faltantes:\n" + "\n".join(f"- {e}" for e in f.elementos_faltantes)
    return texto


def _col6(f: FilaPapel, docs: dict[str, Documento]) -> str:
    a = f.articulo
    norma = docs[a.doc_id].nombre if a.doc_id in docs else a.doc_id
    pag = f"p. {a.pagina_inicio}" if a.pagina_inicio == a.pagina_fin else f"pp. {a.pagina_inicio}-{a.pagina_fin}"
    lineas = [f"Norma: {norma}, {pag}"]
    if _con_respaldo(f):
        r = f.respaldo
        pr = f"p. {r.pagina_inicio}" if r.pagina_inicio == r.pagina_fin else f"pp. {r.pagina_inicio}-{r.pagina_fin}"
        man = docs[r.doc_id].nombre if r.doc_id in docs else r.doc_id
        lineas.append(f"Manual: {man}, sección {r.identificador}, {pr}")
    if f.par and _con_respaldo(f):
        scores = [s for s in (f.par.score_v1, f.par.score_v2) if s is not None and math.isfinite(s)]
        if scores:
            lineas.append(f"Similitud: {max(scores):.3f}")
    if f.par:                                          # los avisos de revisión nunca se ocultan
        v = f.par.veredicto
        if not f.par.cita_norma_verificada or (v and v.cita_manual and not f.par.cita_manual_verificada):
            lineas.append("Aviso: la cita no se verificó contra el texto fuente")
        if f.par.requiere_revision:
            lineas.append("Requiere revisión del auditor")
    if a.seccionado_incierto:
        lineas.append("Aviso: el límite del artículo es incierto")
    return "\n".join(lineas)


def _hoja1(ws, entrada: EntradaPapel) -> None:
    docs = {d.id: d for d in [entrada.manual, *entrada.normativas]}
    conclusion = entrada.conclusion.strip()
    if conclusion and not entrada.conclusion_editada and not conclusion.startswith(PREFIJO_BORRADOR):
        conclusion = PREFIJO_BORRADOR + "\n" + conclusion
    leyenda = "  ·  ".join(f"{m} = {SIGNIFICADO_MARCA[m]}" for m in "ALRXP")
    valores = ["", f"Evaluación de cumplimiento normativo — {entrada.manual.nombre}", OBJETIVO, "",
               _fuente(entrada), leyenda, conclusion, "", "",
               (entrada.fecha_ejecucion or date.today()).strftime("%d/%m/%Y")]
    ancho_valor = sum(ANCHOS_MATRIZ[1:])
    for i, (etiqueta, valor) in enumerate(zip(ETIQUETAS, valores), start=1):
        escribir(ws, i, 1, etiqueta, estilo=estilo_etiqueta)
        escribir(ws, i, 2, valor)
        for col in range(3, 7):                       # bordes de la zona combinada
            ws.cell(row=i, column=col).border = BORDE
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
        ws.row_dimensions[i].height = alto_fila([(valor, ancho_valor), (etiqueta, ANCHOS_MATRIZ[0])])
    h = FILA_ENCABEZADO_MATRIZ
    for col, nombre in enumerate(COLUMNAS, start=1):
        escribir(ws, h, col, nombre, estilo=estilo_encabezado)
        ws.column_dimensions[chr(64 + col)].width = ANCHOS_MATRIZ[col - 1]
    ajustar_impresion(ws, h)
    ws.freeze_panes = ws.cell(row=h + 1, column=1)    # paneles congelados bajo el encabezado de la matriz

    fila = h + 1
    inicio_grupo, clave_grupo = fila, None
    filas = _ordenar(entrada)
    for f in filas:
        a = f.articulo
        celdas = [_col1(f, docs), _col2(f), _col3(f), f.marca, _col5(f), _col6(f, docs)]
        for col, texto in enumerate(celdas, start=1):
            pagina = a.pagina_inicio if col in (1, 2, 5) else (f.respaldo.pagina_inicio if _con_respaldo(f) else None)
            if col == 4:
                escribir(ws, fila, col, texto, estilo=estilo_marca(f.marca))
            else:
                escribir(ws, fila, col, texto, pagina=pagina)
        ws.row_dimensions[fila].height = alto_fila([(celdas[i], ANCHOS_MATRIZ[i]) for i in (1, 2, 4, 5)])
        # Combinar la columna 1 (verticalmente) mientras la fuente sea la misma
        if clave_grupo is not None and celdas[0] != clave_grupo:
            _combinar_col1(ws, inicio_grupo, fila - 1)
            inicio_grupo = fila
        clave_grupo = celdas[0]
        fila += 1
    if filas:
        _combinar_col1(ws, inicio_grupo, fila - 1)


def _combinar_col1(ws, desde: int, hasta: int) -> None:
    if hasta > desde:
        ws.merge_cells(start_row=desde, start_column=1, end_row=hasta, end_column=1)


def generar_papel(entrada: EntradaPapel, destino) -> dict[str, int]:
    """Escribe el Excel en `destino` (ruta o archivo en memoria) y devuelve los conteos por marca."""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Papel de trabajo"
    _hoja1(ws1, entrada)
    annex.escribir_anexo(wb.create_sheet("Anexo técnico"), entrada)
    wb.save(destino)
    return conteos_papel(entrada)
